import base64
import hashlib
import html
import io
import json
import re
import threading
import uuid
import zipfile
from datetime import date, datetime
from pathlib import Path
from urllib.error import HTTPError, URLError
from urllib.request import Request, urlopen

import extract_msg
import msal
import pymupdf
from flask import Flask, jsonify, render_template, request, send_file
from openpyxl import load_workbook


app = Flask(__name__)
app.config["MAX_CONTENT_LENGTH"] = 50 * 1024 * 1024
TEMPLATE_DIR = Path(app.instance_path) / "email_templates"
TEMPLATE_DIR.mkdir(parents=True, exist_ok=True)
workbooks = {}
analyses = {}
auth_sessions = {}


def normalized(value):
    return re.sub(r"\s+", " ", str(value or "")).strip()


def excel_value(value):
    if value is None:
        return ""
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    return normalized(value)


def detect_work(text):
    text = re.sub(r"\s+", " ", text).lower()
    no_work_phrases = (
        "no remedial works required", "no works required", "no action is required",
        "no further action required", "no recommendations", "no significant findings",
        "no deficiencies identified",
    )
    has_no_work_statement = any(phrase in text for phrase in no_work_phrases)
    work_patterns = (
        r"\bworks?\s+(?:are\s+)?required\b",
        r"\b(?:remedial|corrective)\s+works?\b",
        r"\baction\s+(?:required|plan)\b",
        r"\brecommendations?\s+(?:for|to)\s+(?:action|remedial|improvement)",
        r"\bpriority\s*[1-3]\b",
        r"\bcategory\s*[1-3]\s+(?:deficiency|finding|action)",
        r"\bdeficiencies\s+(?:were\s+)?identified\b",
    )
    pattern_text = text
    if has_no_work_statement:
        for phrase in no_work_phrases:
            pattern_text = pattern_text.replace(phrase, " ")
    for pattern in work_patterns:
        match = re.search(pattern, pattern_text)
        if match:
            if has_no_work_statement:
                return None, "Report contains conflicting works and no-works language; review manually."
            start = max(0, match.start() - 90)
            end = min(len(pattern_text), match.end() + 150)
            return True, pattern_text[start:end].strip()
    if has_no_work_statement:
        return False, "Report states that no works or actions are required."
    return None, "No clear works status detected; review the report."


def pdf_from_embedded(data):
    start = data.find(b"%PDF-")
    if start < 0:
        return None
    end = data.rfind(b"%%EOF", start)
    if end < 0:
        return None
    pdf_bytes = data[start:end + 5]
    try:
        document = pymupdf.open(stream=pdf_bytes, filetype="pdf")
        text = "\n".join(page.get_text() for page in document)
        document.close()
    except (ValueError, RuntimeError):
        return None
    return pdf_bytes, text


def embedded_pdfs(workbook_bytes):
    found = []
    seen = set()
    try:
        archive = zipfile.ZipFile(io.BytesIO(workbook_bytes))
    except zipfile.BadZipFile as error:
        raise ValueError("Upload an .xlsx or .xlsm workbook.") from error
    with archive:
        for item in archive.infolist():
            if "embedding" not in item.filename.lower() and not item.filename.lower().endswith(".pdf"):
                continue
            raw = archive.read(item)
            candidate = pdf_from_embedded(raw)
            if not candidate:
                continue
            pdf_bytes, text = candidate
            digest = hashlib.sha256(pdf_bytes).hexdigest()
            if digest in seen:
                continue
            seen.add(digest)
            names = re.findall(rb"[A-Za-z0-9 _().-]{1,120}\.pdf", raw, re.IGNORECASE)
            filename = names[-1].decode("ascii", errors="ignore").strip() if names else ""
            if not filename:
                filename = f"Embedded FRA {len(found) + 1}.pdf"
            has_work, evidence = detect_work(text)
            found.append({
                "id": uuid.uuid4().hex,
                "name": Path(filename).name,
                "data": base64.b64encode(pdf_bytes).decode("ascii"),
                "has_work": has_work,
                "evidence": evidence,
            })
    return found


def load_template(template_id):
    if not re.fullmatch(r"[a-f0-9]{32}", str(template_id or "")):
        return None
    path = TEMPLATE_DIR / f"{template_id}.json"
    if not path.exists():
        return None
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError):
        return None


def list_templates():
    templates = []
    for path in TEMPLATE_DIR.glob("*.json"):
        try:
            item = json.loads(path.read_text(encoding="utf-8"))
            templates.append({"id": item["id"], "name": item["name"], "subject": item["subject"]})
        except (OSError, KeyError, json.JSONDecodeError):
            continue
    return sorted(templates, key=lambda item: item["name"].casefold())


@app.get("/")
@app.get("/koritsu")
def index():
    return render_template("index.html")


@app.get("/koritsu/download-simulation")
def download_simulation():
    workbook_path = Path(app.root_path) / "Koritsu_simulation.xlsx"
    if not workbook_path.is_file():
        return jsonify(error="The simulation workbook is not available."), 404
    return send_file(
        workbook_path,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True,
        download_name="Koritsu_simulation.xlsx",
    )


@app.get("/api/templates")
def templates():
    return jsonify(templates=list_templates())


@app.get("/api/templates/<template_id>")
def template_detail(template_id):
    template = load_template(template_id)
    if not template:
        return jsonify(error="Template not found."), 404
    return jsonify(template)


@app.post("/api/templates")
def import_templates():
    uploads = request.files.getlist("templates")
    if not uploads:
        return jsonify(error="Choose one or more Outlook .oft template files."), 400
    imported = []
    errors = []
    for upload in uploads:
        filename = Path(upload.filename or "").name
        if Path(filename).suffix.lower() != ".oft":
            errors.append(f"{filename or 'File'}: only .oft files are supported.")
            continue
        temporary_path = None
        try:
            contents = upload.read()
            if not contents:
                raise ValueError("The file is empty.")
            temporary_path = TEMPLATE_DIR / f"upload-{uuid.uuid4().hex}.oft"
            temporary_path.write_bytes(contents)
            message = extract_msg.Message(str(temporary_path))
            subject = normalized(message.subject) or Path(filename).stem
            html_body = getattr(message, "htmlBody", None)
            plain_body = getattr(message, "body", None) or ""
            if isinstance(html_body, bytes):
                html_body = html_body.decode("utf-8", errors="replace")
            if isinstance(plain_body, bytes):
                plain_body = plain_body.decode("utf-8", errors="replace")
            content = str(html_body).strip() if html_body else str(plain_body)
            content_type = "html" if html_body else "text"
            if not content.strip():
                raise ValueError("No message body was found in this template.")
            template_id = uuid.uuid4().hex
            item = {"id": template_id, "name": Path(filename).stem, "subject": subject,
                    "body": content, "content_type": content_type}
            (TEMPLATE_DIR / f"{template_id}.json").write_text(json.dumps(item), encoding="utf-8")
            close = getattr(message, "close", None)
            if close:
                close()
            imported.append({"id": template_id, "name": item["name"], "subject": subject})
        except Exception as error:
            errors.append(f"{filename}: {error}")
        finally:
            if temporary_path and temporary_path.exists():
                temporary_path.unlink()
    return jsonify(imported=imported, errors=errors, templates=list_templates())


@app.delete("/api/templates/<template_id>")
def delete_template(template_id):
    if not re.fullmatch(r"[a-f0-9]{32}", template_id):
        return jsonify(error="Invalid template id."), 400
    path = TEMPLATE_DIR / f"{template_id}.json"
    if not path.exists():
        return jsonify(error="Template not found."), 404
    path.unlink()
    return jsonify(templates=list_templates())


@app.post("/api/spreadsheet")
def spreadsheet():
    upload = request.files.get("spreadsheet")
    if not upload or not upload.filename:
        return jsonify(error="Choose an Excel spreadsheet to continue."), 400
    try:
        workbook_bytes = upload.read()
        workbook = load_workbook(io.BytesIO(workbook_bytes), read_only=True, data_only=True)
        sheet = workbook.active
        values = sheet.iter_rows(values_only=True)
        headers = [normalized(value).casefold() for value in next(values, [])]
        column_aliases = {
            "managing_agent": ("managing agent",),
            "email": ("ma email", "managing agent email"),
            "property": ("property address",),
            "postcode": ("postcode", "post code"),
            "report": ("fra report", "fra report.pdf"),
            "inspection_date": ("inspection date",),
            "due_date": ("due date",),
            "notes": ("notes",),
        }
        columns = {}
        for key, aliases in column_aliases.items():
            columns[key] = next((headers.index(alias) for alias in aliases if alias in headers), None)
        required = ("managing_agent", "email", "property", "postcode", "report")
        missing = [key.replace("_", " ") for key in required if columns[key] is None]
        if missing:
            workbook.close()
            return jsonify(error="Missing required headings: " + ", ".join(missing)), 400
        records = []
        for row_number, row in enumerate(values, start=2):
            converted = [excel_value(value) for value in row]
            if not any(converted):
                continue
            records.append({
                "row_number": row_number,
                "managing_agent": converted[columns["managing_agent"]] if columns["managing_agent"] < len(converted) else "",
                "recipient": converted[columns["email"]] if columns["email"] < len(converted) else "",
                "property": converted[columns["property"]] if columns["property"] < len(converted) else "",
                "postcode": converted[columns["postcode"]] if columns["postcode"] < len(converted) else "",
                "report_cell": converted[columns["report"]] if columns["report"] < len(converted) else "",
                "inspection_date": converted[columns["inspection_date"]] if columns["inspection_date"] is not None and columns["inspection_date"] < len(converted) else "",
                "due_date": converted[columns["due_date"]] if columns["due_date"] is not None and columns["due_date"] < len(converted) else "",
                "notes": converted[columns["notes"]] if columns["notes"] is not None and columns["notes"] < len(converted) else "",
            })
        workbook.close()
        if not records:
            return jsonify(error="The first worksheet has no property rows."), 400
        reports = embedded_pdfs(workbook_bytes)
    except Exception as error:
        return jsonify(error=f"Could not read this workbook: {error}"), 400
    upload_id = uuid.uuid4().hex
    workbooks[upload_id] = {"rows": records, "reports": reports}
    return jsonify(
        upload_id=upload_id,
        rows=records,
        reports=[{key: value for key, value in report.items() if key != "data"} for report in reports],
        row_count=len(records),
    )


@app.post("/api/analyze")
def analyze():
    payload = request.get_json(force=True)
    workbook = workbooks.get(payload.get("upload_id"))
    if not workbook:
        return jsonify(error="That workbook session has expired. Upload it again."), 400
    report_map = payload.get("report_map", {})
    report_by_id = {report["id"]: report for report in workbook["reports"]}
    results = []
    for row in workbook["rows"]:
        report = report_by_id.get(report_map.get(str(row["row_number"])))
        if not report and row["report_cell"]:
            needle = row["report_cell"].casefold()
            report = next((item for item in workbook["reports"] if needle in item["name"].casefold()), None)
        results.append({**row, "report": report})
    analysis_id = uuid.uuid4().hex
    analyses[analysis_id] = results
    return jsonify(analysis_id=analysis_id, results=results)


@app.post("/api/auth/start")
def auth_start():
    payload = request.get_json(force=True)
    client_id = normalized(payload.get("client_id"))
    tenant_id = normalized(payload.get("tenant_id")) or "organizations"
    if not client_id:
        return jsonify(error="Enter the Microsoft Entra app (client) ID."), 400
    try:
        client = msal.PublicClientApplication(
            client_id, authority=f"https://login.microsoftonline.com/{tenant_id}"
        )
        flow = client.initiate_device_flow(scopes=["Mail.Send"])
        if not flow.get("user_code"):
            return jsonify(error=flow.get("error_description", "Could not start Microsoft sign-in.")), 400
    except (ValueError, RuntimeError) as error:
        return jsonify(error=f"Could not start Microsoft sign-in: {error}"), 400
    auth_id = uuid.uuid4().hex
    auth_sessions[auth_id] = {"status": "pending", "message": flow["message"], "token": None}

    def finish_sign_in():
        try:
            result = client.acquire_token_by_device_flow(flow)
            if "access_token" in result:
                auth_sessions[auth_id].update({"status": "connected", "token": result["access_token"]})
            else:
                auth_sessions[auth_id].update({
                    "status": "error",
                    "error": result.get("error_description", "Microsoft sign-in was not completed."),
                })
        except Exception as error:
            auth_sessions[auth_id].update({"status": "error", "error": str(error)})

    threading.Thread(target=finish_sign_in, daemon=True).start()
    return jsonify(auth_id=auth_id, message=flow["message"])


@app.get("/api/auth/status/<auth_id>")
def auth_status(auth_id):
    session = auth_sessions.get(auth_id)
    if not session:
        return jsonify(error="Sign-in session not found."), 404
    return jsonify(status=session["status"], message=session.get("message"), error=session.get("error"))


@app.post("/api/auth/disconnect")
def auth_disconnect():
    payload = request.get_json(force=True)
    auth_id = payload.get("auth_id")
    session = auth_sessions.get(auth_id)
    if session and session.get("status") == "connected":
        auth_sessions.pop(auth_id, None)
    return jsonify(status="disconnected")


def template_fields(row):
    return {
        "managing agent": row["managing_agent"], "agent": row["managing_agent"],
        "ma email": row["recipient"], "email": row["recipient"], "recipient": row["recipient"],
        "property address": row["property"], "property": row["property"],
        "postcode": row["postcode"], "inspection date": row["inspection_date"],
        "due date": row["due_date"], "notes": row["notes"],
    }


def render_template_text(text, fields, is_html=False):
    pattern = re.compile(r"\{\{\s*([^{}]+?)\s*\}\}|«([^»]+)»|\[\[\s*([^\]]+?)\s*\]\]", re.IGNORECASE)

    def replace(match):
        key = next(group for group in match.groups() if group is not None).strip().casefold()
        value = fields.get(key)
        if value is None:
            return match.group(0)
        value = str(value)
        return html.escape(value) if is_html else value

    return pattern.sub(replace, text)


def graph_json_request(path, token, payload=None, method="POST"):
    body = json.dumps(payload).encode("utf-8") if payload is not None else None
    headers = {"Authorization": f"Bearer {token}"}
    if body is not None:
        headers["Content-Type"] = "application/json"
    graph_request = Request(f"https://graph.microsoft.com/v1.0{path}", data=body, headers=headers, method=method)
    with urlopen(graph_request, timeout=45) as response:
        response_data = response.read()
        return response.status, json.loads(response_data) if response_data else {}


def send_graph_message(token, message, report, pdf_bytes):
    if len(pdf_bytes) <= 2 * 1024 * 1024:
        message["attachments"] = [{
            "@odata.type": "#microsoft.graph.fileAttachment",
            "name": report["name"],
            "contentType": "application/pdf",
            "contentBytes": base64.b64encode(pdf_bytes).decode("ascii"),
        }]
        return graph_json_request("/me/sendMail", token, {"message": message, "saveToSentItems": True})

    _, draft = graph_json_request("/me/messages", token, message)
    message_id = draft["id"]
    _, session = graph_json_request(
        f"/me/messages/{message_id}/attachments/createUploadSession", token,
        {"AttachmentItem": {
            "attachmentType": "file", "name": report["name"],
            "size": len(pdf_bytes), "contentType": "application/pdf",
        }},
    )
    upload_url = session["uploadUrl"]
    chunk_size = 320 * 1024
    for start in range(0, len(pdf_bytes), chunk_size):
        chunk = pdf_bytes[start:start + chunk_size]
        end = start + len(chunk) - 1
        chunk_request = Request(
            upload_url, data=chunk,
            headers={
                "Content-Length": str(len(chunk)),
                "Content-Range": f"bytes {start}-{end}/{len(pdf_bytes)}",
                "Content-Type": "application/octet-stream",
            },
            method="PUT",
        )
        with urlopen(chunk_request, timeout=60):
            pass
    return graph_json_request(f"/me/messages/{message_id}/send", token, method="POST")


@app.post("/api/send")
def send_emails():
    payload = request.get_json(force=True)
    rows = analyses.get(payload.get("analysis_id"))
    if rows is None:
        return jsonify(error="That review session has expired. Import the workbook again."), 400
    auth = auth_sessions.get(payload.get("auth_id"))
    if not auth or auth.get("status") != "connected" or not auth.get("token"):
        return jsonify(error="Connect to Microsoft 365 before sending."), 401
    template = load_template(payload.get("template_id"))
    if not template:
        return jsonify(error="Choose an imported Outlook template."), 400
    selected_rows = set(payload.get("row_numbers", []))
    selected = [row for row in rows if row["row_number"] in selected_rows]
    statuses = payload.get("reviewed_statuses", {})
    if not selected:
        return jsonify(error="Select at least one property marked works found."), 400
    if any(statuses.get(str(row["row_number"])) != "works" for row in selected):
        return jsonify(error="Only properties reviewed as works found can be sent."), 400
    if any(not row["report"] or not row["recipient"] for row in selected):
        return jsonify(error="Every selected property needs an assigned PDF and managing-agent email."), 400

    sent = []
    failures = []
    for row in selected:
        report = row["report"]
        pdf_bytes = base64.b64decode(report["data"])
        fields = template_fields(row)
        content_is_html = template["content_type"] == "html"
        message = {
            "subject": render_template_text(template["subject"], fields),
            "body": {
                "contentType": "HTML" if content_is_html else "Text",
                "content": render_template_text(template["body"], fields, is_html=content_is_html),
            },
            "toRecipients": [{"emailAddress": {"address": row["recipient"]}}],
        }
        try:
            response_status, _ = send_graph_message(auth["token"], message, report, pdf_bytes)
            if response_status in (200, 201, 202):
                sent.append(row["row_number"])
            else:
                failures.append({"row_number": row["row_number"], "error": f"Microsoft Graph returned HTTP {response_status}."})
        except HTTPError as error:
            detail = error.read().decode("utf-8", errors="replace")
            failures.append({"row_number": row["row_number"], "error": f"Microsoft Graph HTTP {error.code}: {detail[:400]}"})
        except (URLError, TimeoutError, OSError) as error:
            failures.append({"row_number": row["row_number"], "error": str(error)})
    return jsonify(sent=sent, failures=failures)


if __name__ == "__main__":
    app.run(host="127.0.0.1", port=5000, debug=False, threaded=True)